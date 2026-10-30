"""Resume collector: PDF (pypdf), DOCX (python-docx), TXT and MD.

The text is redacted and split at paragraph breaks into parts small enough
for one extraction call each. PDF and DOCX need the ``onboard`` extra.
"""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any

from .. import config
from .items import FILE, UNCHANGED, Collected, SourceItem, file_hash, seen
from .privacy import redact

SUFFIXES = (".pdf", ".docx", ".txt", ".md")
PART_CHARS = 6000
MAX_RESUME_BYTES = 5_000_000
EXTRA_HINT = 'reading PDF and DOCX needs the onboard extra: pip install -e ".[onboard]"'


def _pdf_text(path: Path) -> str:
    try:
        from pypdf import PdfReader  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - the extra is in the dev lockfile
        raise RuntimeError(EXTRA_HINT) from exc
    return "\n\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def _docx_text(path: Path) -> str:
    try:
        import docx  # noqa: PLC0415
    except ImportError as exc:  # pragma: no cover - the extra is in the dev lockfile
        raise RuntimeError(EXTRA_HINT) from exc
    document = docx.Document(str(path))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append(" | ".join(cell.text for cell in row.cells))
    return "\n".join(lines)


def read_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix not in SUFFIXES:
        raise ValueError(f"resume must be one of {', '.join(SUFFIXES)}: {path.name}")
    if path.stat().st_size > MAX_RESUME_BYTES:
        raise ValueError(f"resume is over {MAX_RESUME_BYTES // 1_000_000} MB: {path.name}")
    if suffix == ".pdf":
        return _pdf_text(path)
    if suffix == ".docx":
        return _docx_text(path)
    return path.read_text(encoding="utf-8", errors="replace")


def split_parts(text: str, limit: int = PART_CHARS) -> list[str]:
    """Paragraph aligned parts of at most ``limit`` characters."""
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]
    parts: list[str] = []
    current = ""
    for whole in paragraphs:
        para = whole
        while len(para) > limit:  # one huge paragraph: cut it
            parts.extend([current] if current else [])
            current = ""
            parts.append(para[:limit])
            para = para[limit:]
        if current and len(current) + len(para) + 2 > limit:
            parts.append(current)
            current = ""
        current = f"{current}\n\n{para}" if current else para
    if current:
        parts.append(current)
    return parts


def collect(source: dict[str, Any]) -> Collected:
    path = Path(source["value"]).expanduser()
    out = Collected()
    try:
        out.stamps[FILE] = file_hash(path)
        if seen(source, FILE) == out.stamps[FILE]:
            out.notes.append(UNCHANGED)
            return out
        text = read_text(path)
    except (OSError, ValueError, RuntimeError) as exc:
        out.notes.append(f"resume not read: {exc}")
        return out
    clean = re.sub(r"[ \t]+", " ", redact(text)).strip()
    if not clean:
        out.notes.append(f"resume has no text (a scanned image?): {path.name}")
        return out
    out.read.append(path.name)
    for n, part in enumerate(split_parts(clean, min(PART_CHARS, config.ONBOARD_BATCH_CHARS))):
        out.items.append(SourceItem(source["id"], "resume", path.name, part, n))
    return out
