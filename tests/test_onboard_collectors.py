"""Onboarding collectors and the privacy rules, on made up fixtures only."""

from __future__ import annotations

import json
import os
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock

import docx
import pytest
import requests

from sourcer import config
from sourcer.onboard import collect, folders, github, linkedin, privacy, resume, store
from sourcer.onboard.items import source_spec

# made up secrets, built in pieces so no scanner mistakes them for real ones
FAKE_KEY = "sk-" + "ant-" + "api03-" + "Zx9" * 8
FAKE_GH = "ghp_" + "Q1w2E3r4" * 4
FAKE_EMAIL = "pat.example" + "@" + "example.org"


def make_pdf(path: Path, lines: list[str]) -> Path:
    """A one page PDF with real text objects, written by hand."""
    body = " ".join(f"({line}) Tj T*" for line in lines)
    content = f"BT /F1 12 Tf 72 720 Td 14 TL {body} ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        "/Resources << /Font << /F1 5 0 R >> >> >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for n, obj in enumerate(objects, 1):
        offsets.append(len(out))
        out += f"{n} 0 obj\n{obj}\nendobj\n".encode()
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    out += b"".join(f"{o:010d} 00000 n \n".encode() for o in offsets)
    out += (
        f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    )
    path.write_bytes(out)
    return path


def spec(kind: str, value: str, **extra: Any) -> dict[str, Any]:
    return source_spec(kind, value, **extra)


# ---------------------------------------------------------------- resume


def test_resume_txt_is_redacted_and_split(tmp_path: Path) -> None:
    text = f"Alex Example\n{FAKE_EMAIL} | 555-123-4567\n\n" + "\n\n".join(
        f"Built pipeline number {n} in Python and SQL." for n in range(400)
    )
    path = tmp_path / "cv.txt"
    path.write_text(text)
    got = resume.collect(spec("resume", str(path)))
    assert got.read == ["cv.txt"]
    assert len(got.items) > 1
    joined = "\n".join(i.text for i in got.items)
    assert FAKE_EMAIL not in joined and "555-123-4567" not in joined
    assert "[email]" in joined and "[phone]" in joined
    assert all(len(i.text) <= resume.PART_CHARS for i in got.items)
    ids = [i.id for i in got.items]
    assert len(set(ids)) == len(ids)
    assert got.items[0].id == resume.collect(spec("resume", str(path))).items[0].id


def test_resume_pdf(tmp_path: Path) -> None:
    path = make_pdf(
        tmp_path / "cv.pdf", ["Data Analyst Intern, Example Corp", "Built dashboards in Tableau"]
    )
    got = resume.collect(spec("resume", str(path)))
    assert "Built dashboards in Tableau" in got.items[0].text
    assert got.items[0].source_type == "resume" and got.items[0].location == "cv.pdf"


def test_resume_docx(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_paragraph("Research Assistant, Example University")
    document.add_paragraph("Cleaned survey data with pandas.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Skills"
    table.rows[0].cells[1].text = "R, Stata"
    path = tmp_path / "cv.docx"
    document.save(str(path))
    text = resume.collect(spec("resume", str(path))).items[0].text
    assert "Cleaned survey data with pandas." in text and "Skills | R, Stata" in text


def test_resume_problems_become_notes(tmp_path: Path) -> None:
    odd = tmp_path / "cv.rtf"
    odd.write_text("x")
    assert "must be one of" in resume.collect(spec("resume", str(odd))).notes[0]
    assert "not read" in resume.collect(spec("resume", str(tmp_path / "gone.pdf"))).notes[0]
    blank = tmp_path / "blank.md"
    blank.write_text("   \n")
    assert "no text" in resume.collect(spec("resume", str(blank))).notes[0]


def test_resume_too_big(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(resume, "MAX_RESUME_BYTES", 3)
    path = tmp_path / "cv.md"
    path.write_text("long enough")
    assert "over" in resume.collect(spec("resume", str(path))).notes[0]


def test_split_parts_cuts_a_huge_paragraph() -> None:
    parts = resume.split_parts("short\n\n" + "x" * 25 + "\n\nend", limit=10)
    assert parts[0] == "short" and all(len(p) <= 10 for p in parts) and parts[-1].endswith("end")


# ---------------------------------------------------------------- github


class FakeResponse:
    def __init__(self, status: int = 200, body: Any = None, headers: dict[str, str] | None = None):
        self.status_code = status
        self._body = body
        self.headers = headers or {}
        self.text = body if isinstance(body, str) else json.dumps(body)
        self.content = self.text.encode()

    def json(self) -> Any:
        return self._body

    def raise_for_status(self) -> None:
        if self.status_code >= 400:
            raise requests.HTTPError(f"{self.status_code}")


REPOS = [
    {"name": "churn-model", "description": "Predicts churn", "topics": ["ml"], "language": "Python",
     "created_at": "2024-01-02T00:00:00Z", "pushed_at": "2025-03-04T00:00:00Z",
     "stargazers_count": 3, "html_url": "https://github.com/octo-example/churn-model",
     "fork": False, "archived": True, "homepage": "https://example.org"},
    {"name": "forked-lib", "fork": True},
    {"name": "notes", "description": None, "topics": [], "language": None, "fork": False},
]  # fmt: skip


def fake_github(monkeypatch: pytest.MonkeyPatch, routes: dict[str, FakeResponse]) -> MagicMock:
    calls = MagicMock()

    def get(url: str, **kwargs: Any) -> FakeResponse:
        calls(url, **kwargs)
        path = url.removeprefix(github.API).split("?")[0]
        return routes.get(path, FakeResponse(404, {}))

    monkeypatch.setattr(github.net, "get", get)
    return calls


def test_github_reads_public_repos(onboard_paths: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    calls = fake_github(
        monkeypatch,
        {
            "/users/octo-example/repos": FakeResponse(200, REPOS),
            "/repos/octo-example/churn-model/languages":
                FakeResponse(200, {"Python": 900, "Shell": 100}),
            "/repos/octo-example/churn-model/readme":
                FakeResponse(200, f"# Churn\nUses scikit-learn. {FAKE_GH}"),
        },
    )  # fmt: skip
    got = github.collect(spec("github", "@octo-example"))
    assert got.read == ["churn-model", "notes"]  # the fork is skipped
    first = got.items[0]
    assert first.location == "https://github.com/octo-example/churn-model"
    assert "Languages: Python 90%, Shell 10%" in first.text
    assert "Last pushed: 2025-03-04" in first.text and "Archived: yes" in first.text
    assert "scikit-learn" in first.text and FAKE_GH not in first.text
    assert got.items[1].location == "https://github.com/octo-example/notes"
    # unauthenticated, and every call is logged without content
    assert all("Authorization" not in c.kwargs["headers"] for c in calls.call_args_list)
    log = store.read_audit()
    assert {e["destination"] for e in log} == {"api.github.com"}
    assert all(set(e) == {"ts", "destination", "action", "sources", "chars"} for e in log)


def test_github_token_is_sent_when_set(
    onboard_paths: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setenv("GITHUB_TOKEN", "test-token-value")
    calls = fake_github(monkeypatch, {"/users/octo-example/repos": FakeResponse(200, [])})
    github.collect(spec("github", "octo-example"))
    assert calls.call_args.kwargs["headers"]["Authorization"] == "Bearer test-token-value"
    assert calls.call_args.kwargs["allow_redirects"] is False


def test_github_rate_limit_keeps_what_was_read(
    onboard_paths: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    limited = FakeResponse(
        403, {}, {"X-RateLimit-Remaining": "0", "X-RateLimit-Reset": "1700000000"}
    )
    fake_github(
        monkeypatch,
        {
            "/users/octo-example/repos": FakeResponse(200, REPOS),
            "/repos/octo-example/churn-model/languages": FakeResponse(200, {}),
            "/repos/octo-example/churn-model/readme": FakeResponse(301, {}, {"Location": "x"}),
            "/repos/octo-example/notes/languages": limited,
        },
    )
    got = github.collect(spec("github", "octo-example"))
    assert len(got.items) == 1 and "README" not in got.items[0].text  # a redirect is no README
    assert "rate limit" in got.notes[0] and "GITHUB_TOKEN" in got.notes[0]


def test_github_errors(onboard_paths: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    fake_github(monkeypatch, {})
    assert "not found" in github.collect(spec("github", "nobody-here")).notes[0]
    assert "not a GitHub username" in github.collect(spec("github", "bad name!")).notes[0]
    fake_github(monkeypatch, {"/users/octo-example/repos": FakeResponse(500, {})})
    assert "not read" in github.collect(spec("github", "octo-example")).notes[0]
    fake_github(monkeypatch, {"/users/octo-example/repos": FakeResponse(429, {})})
    assert "rate limit" in github.collect(spec("github", "octo-example")).notes[0]


# ---------------------------------------------------------------- folders


def git(root: Path, *args: str, date: str = "2024-02-03T10:00:00") -> None:
    env = {**os.environ, "GIT_AUTHOR_DATE": date, "GIT_COMMITTER_DATE": date,
           "GIT_AUTHOR_NAME": "Test", "GIT_AUTHOR_EMAIL": "test@example.org",
           "GIT_COMMITTER_NAME": "Test", "GIT_COMMITTER_EMAIL": "test@example.org",
           "GIT_CONFIG_GLOBAL": os.devnull, "GIT_CONFIG_NOSYSTEM": "1"}  # fmt: skip
    subprocess.run(["git", "-C", str(root), *args], check=True, capture_output=True, env=env)


def make_project(root: Path) -> Path:
    root.mkdir(parents=True)
    (root / "README.md").write_text(
        f"# Thesis\nForecasting demand with PyTorch.\nContact {FAKE_EMAIL}\nkey: {FAKE_KEY}\n"
    )
    (root / "pyproject.toml").write_text('[project]\nname = "thesis"\ndependencies = ["torch"]\n')
    (root / "train.py").write_text("print('train')\n")
    (root / "src").mkdir()
    (root / "src" / "model.py").write_text("x = 1\n")
    (root / ".env").write_text(f"ANTHROPIC_API_KEY={FAKE_KEY}\n")
    (root / "id_rsa").write_text("-----BEGIN OPENSSH PRIVATE KEY-----\nabc\n")
    (root / "secrets.json").write_text("{}")
    (root / "server.pem").write_text("pem")
    (root / "node_modules" / "lib").mkdir(parents=True)
    (root / "node_modules" / "lib" / "index.js").write_text("module.exports = 1")
    (root / ".gitignore").write_text("private_notes/\n*.log\n!keep.log\n/scratch.txt\n")
    (root / "private_notes").mkdir()
    (root / "private_notes" / "diary.md").write_text("private")
    (root / "run.log").write_text("log")
    (root / "keep.log").write_text("log")
    (root / "scratch.txt").write_text("scratch")
    (root / "data.bin").write_bytes(b"\0\1\2")
    (root / "big.csv").write_text("x" * 50)
    return root


@pytest.fixture
def project(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> Path:
    monkeypatch.setattr(config, "ONBOARD_MAX_FILE_BYTES", 40)
    return make_project(tmp_path / "thesis")


def test_folder_reads_only_structure_and_skips_private_files(
    project: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    opened: list[str] = []
    real_read_text, real_open = Path.read_text, Path.open

    def spy_read_text(self: Path, *a: Any, **k: Any) -> str:
        opened.append(self.name)
        return real_read_text(self, *a, **k)

    def spy_open(self: Path, *a: Any, **k: Any) -> Any:
        opened.append(self.name)
        return real_open(self, *a, **k)

    monkeypatch.setattr(Path, "read_text", spy_read_text)
    monkeypatch.setattr(Path, "open", spy_open)
    monkeypatch.setattr(config, "ONBOARD_MAX_FILE_BYTES", 400)
    got = folders.collect(spec("folder", str(project)))
    never = {".env", "id_rsa", "secrets.json", "server.pem", "index.js", "diary.md", "run.log",
             "keep.log", "scratch.txt"}  # fmt: skip
    assert not never & set(opened), never & set(opened)
    assert sorted(got.read) == ["README.md", "pyproject.toml"]
    text = got.items[0].text
    assert got.items[0].location == "."
    assert "Forecasting demand with PyTorch." in text and 'dependencies = ["torch"]' in text
    assert ".py 2" in text
    assert FAKE_KEY not in text and FAKE_EMAIL not in text
    assert got.skipped["secrets"] >= 4 and got.skipped["ignored"] >= 3
    assert got.skipped["dependencies"] == 1


def test_folder_size_and_binary_rules(project: Path) -> None:
    found = folders.plan(project)
    assert found.skipped["large files"] >= 1 and found.skipped["binaries"] == 1


@pytest.mark.skipif(shutil.which("git") is None, reason="needs git")
def test_folder_git_dates(project: Path) -> None:
    git(project, "init", "-q")
    git(project, "add", "README.md")
    git(project, "commit", "-qm", "first", date="2023-09-01T10:00:00")
    git(project, "add", "train.py")
    git(project, "commit", "-qm", "second", date="2024-05-06T10:00:00")
    text = folders.collect(spec("folder", str(project))).items[0].text
    assert "First commit: 2023-09-01" in text and "Last commit: 2024-05-06" in text
    assert "Commits: 2" in text


def test_folder_of_projects_with_excludes(tmp_path: Path) -> None:
    root = tmp_path / "code"
    make_project(root / "thesis")
    (root / "old").mkdir()
    (root / "old" / "README.md").write_text("# Old\n")
    (root / "node_modules").mkdir()
    (root / "loose-notes").mkdir()
    (root / "file.txt").write_text("not a project")
    got = folders.collect(spec("folder", str(root), excludes=["old"]))
    assert [i.location for i in got.items] == ["thesis"]
    assert "thesis/README.md" in got.read
    both = folders.plan(root)
    assert [p.location for p in both.projects] == ["old", "thesis"]


def test_folder_links_are_never_followed(tmp_path: Path) -> None:
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "notes.md").write_text("outside")
    root = tmp_path / "proj"
    root.mkdir()
    (root / "README.md").write_text("# Proj\n")
    (root / "linked").symlink_to(outside, target_is_directory=True)
    (root / "README.txt").symlink_to(outside / "notes.md")
    found = folders.plan(root)
    assert found.files == ["README.md"]
    assert found.skipped["links"] == 2


def test_folder_missing_or_empty(tmp_path: Path) -> None:
    assert "not found" in folders.collect(spec("folder", str(tmp_path / "nope"))).notes[0]
    empty = tmp_path / "empty"
    empty.mkdir()
    assert "no projects" in folders.collect(spec("folder", str(empty))).notes[0]
    assert folders.git_dates(empty) is None


# ---------------------------------------------------------------- linkedin


def make_linkedin_zip(path: Path) -> Path:
    with zipfile.ZipFile(path, "w") as z:
        z.writestr("Positions.csv",
                   "Company Name,Title,Description,Location,Started On,Finished On\n"
                   "Example Corp,Data Intern,Built a forecasting model in Python,New York,"
                   "Jun 2024,Aug 2024\n")  # fmt: skip
        z.writestr("Projects.csv", "Title,Description,Url,Started On,Finished On\n"
                   "Bike share study,Mapped rides with GeoPandas,,Jan 2024,\n")  # fmt: skip
        z.writestr("Skills.csv", "Name\nPython\nSQL\n")
        z.writestr("Certifications.csv", "Notes:\nExported data\n\n"
                   "Name,Url,Authority,Started On,Finished On,License Number\n"
                   "Cloud Practitioner,,Example Cloud,Mar 2025,,\n")  # fmt: skip
        z.writestr(
            "Connections.csv", "First Name,Last Name,Email Address\nSam,Sample,sam@example.org\n"
        )
        z.writestr("Education.csv", "")
    return path


def test_linkedin_export(tmp_path: Path) -> None:
    got = linkedin.collect(spec("linkedin", str(make_linkedin_zip(tmp_path / "export.zip"))))
    assert "Connections.csv" not in got.read
    texts = {i.location: i.text for i in got.items}
    assert "Title: Data Intern\nCompany Name: Example Corp" in texts["Positions.csv"]
    assert "Started On: Jun 2024" in texts["Positions.csv"]
    assert "GeoPandas" in texts["Projects.csv"]
    assert texts["Skills.csv"] == "Skills listed on LinkedIn: Python, SQL"
    assert "Authority: Example Cloud" in texts["Certifications.csv"]
    assert all("sam@example.org" not in t for t in texts.values())


def test_linkedin_problems(tmp_path: Path, monkeypatch: pytest.MonkeyPatch) -> None:
    bad = tmp_path / "bad.zip"
    bad.write_text("not a zip")
    assert "not read" in linkedin.collect(spec("linkedin", str(bad))).notes[0]
    other = tmp_path / "other.zip"
    with zipfile.ZipFile(other, "w") as z:
        z.writestr("Connections.csv", "a,b\n")
    assert "no Positions" in linkedin.collect(spec("linkedin", str(other))).notes[0]
    monkeypatch.setattr(linkedin, "MAX_MEMBER_BYTES", 10)
    got = linkedin.collect(spec("linkedin", str(make_linkedin_zip(tmp_path / "e.zip"))))
    assert got.skipped["large files"] >= 3


# ---------------------------------------------------------------- privacy and registry


def test_redact_masks_secrets_and_contact_details() -> None:
    text = (
        f"key {FAKE_KEY} gh {FAKE_GH} aws AKIA{'A' * 16} "
        "password = hunter22hunter jwt eyJabcdefgh.eyJabcdefgh.sigsigsigsig "
        "url https://user:pa55word@example.org/x Bearer abcdefghijklmnopqrstuv "
        f"mail {FAKE_EMAIL} phone (212) 555-0199 +1 212.555.0199\n"
        "-----BEGIN RSA PRIVATE KEY-----\nMIIabc\n-----END RSA PRIVATE KEY-----\n"
        "kept: 2021-2023, 3.14, version 1.2.3"
    )
    out = privacy.redact(text)
    for secret in (FAKE_KEY, FAKE_GH, "AKIA", "hunter22", "eyJ", "pa55word", "abcdefghijklmnop",
                   FAKE_EMAIL, "555-0199", "MIIabc"):  # fmt: skip
        assert secret not in out, secret
    assert "kept: 2021-2023, 3.14, version 1.2.3" in out
    assert privacy.contact_detail(f"write to {FAKE_EMAIL}") and privacy.contact_detail("[phone]")
    assert not privacy.contact_detail("Built a model in Python")


def test_ignore_rules_scope_and_negation(tmp_path: Path) -> None:
    (tmp_path / "sub").mkdir()
    (tmp_path / "sub" / ".gitignore").write_text("# comment\n\n*.csv\n!keep.csv\nout/\n")
    rules = privacy.IgnoreRules(tmp_path, ["docs/*"])
    rules.load(tmp_path)  # no file: nothing added
    rules.load(tmp_path / "sub")
    assert rules.ignored(tmp_path / "sub" / "a.csv")
    assert rules.ignored(tmp_path / "sub" / "keep.csv")  # negations never bring a file back
    assert rules.ignored(tmp_path / "sub" / "out" / "x.py")
    assert not rules.ignored(tmp_path / "a.csv")  # scoped to sub/
    assert rules.ignored(tmp_path / "docs" / "x.md")


def test_skip_reason_unreadable(tmp_path: Path) -> None:
    assert privacy.skip_reason(tmp_path / "missing.txt", 100) == "unreadable"
    assert privacy.dir_skip_reason(".git") == "hidden"
    assert privacy.dir_skip_reason("my-secrets") == "secrets"
    assert privacy.dir_skip_reason("src") is None


def test_registry_and_links(tmp_path: Path) -> None:
    link = spec("link", "https://example.org/portfolio")
    got = collect.collect_all([link, {"id": "x", "type": "unknown", "value": "y"}])
    assert list(got) == [link["id"]]
    assert "not read yet" in got[link["id"]].notes[0]
    with pytest.raises(ValueError):
        source_spec("fax", "1")
    assert spec("github", "a")["id"] == spec("github", "a")["id"] != spec("github", "b")["id"]


def test_git_ignores_programs_named_by_the_repo(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    """A scanned repo's .git/config cannot make git run a program: the command
    line turns off every setting that names one."""
    (tmp_path / ".git").mkdir()
    seen: list[list[str]] = []

    def run(argv: list[str], **kw: Any) -> Any:
        seen.append(argv)
        return MagicMock(returncode=0, stdout="2025-01-02\n2024-03-04\n")

    monkeypatch.setattr(folders.subprocess, "run", run)
    assert folders.git_dates(tmp_path) == ("2024-03-04", "2025-01-02", 2)
    argv = seen[0]
    assert "--no-show-signature" in argv and "--no-pager" in argv
    settings = ("log.showSignature=false", "gpg.program=false", "gpg.ssh.program=false",
                "gpg.x509.program=false", "core.fsmonitor=false", "core.pager=cat")  # fmt: skip
    for setting in settings:
        assert argv[argv.index(setting) - 1] == "-c"
    assert argv.index("log") > argv.index("gpg.program=false")


def test_redact_more_token_formats() -> None:
    samples = {
        '"api_key": "abcdef123456"': "abcdef123456",
        '{"token":"tok_9a8b7c6d5e"}': "tok_9a8b7c6d5e",
        "OPENAI_KEY=abcdefgh1234": "abcdefgh1234",
        "sk_live_" + "a1b2c3d4e5f6g7": "a1b2c3d4e5f6g7",
        "hf_" + "A" * 30: "A" * 30,
        "glpat-" + "x" * 20: "x" * 20,
        "npm_" + "n" * 36: "n" * 36,
        "pypi-" + "p" * 60: "p" * 60,
        "https://hooks.slack.com/services/T000/B000/" + "z" * 24: "z" * 24,
        "postgres://user:pa@ss@db.example.org/x": "pa@ss",
        "call +44 20 7946 0958 now": "7946 0958",
        "or 5551234567 later": "5551234567",
    }
    for text, secret in samples.items():
        assert secret not in privacy.redact(text), text
